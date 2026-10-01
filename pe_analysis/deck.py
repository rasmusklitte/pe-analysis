"""Samler faerdigtegnede 16:9-slides (PNG) i eet PowerPoint-deck paa PPIM-masteren.

Graferne i notebookene er allerede hele slides - kicker, titel, undertitel, salviestreg,
kildelinje og signatur er tegnet med matplotlib. Modulet laegger dem derfor kant-til-kant
paa masterens Blank-layout og tilfoejer kun det, matplotlib ikke kan levere: forside,
agenda med noegletal, sektionsskilleslides, tekstslides og slut-slide med foto og logo.

Brug (fra en analysemappe; masteren ligger i repo-roden og findes automatisk):

    from pe_analysis.deck import byg_deck

    byg_deck("exports/nep_v_tvpi_udvikling.pptx",
             titel="NEP V - TVPI pr. investering", undertitel="Pr. 30.06.2026",
             sektioner=[{"titel": "Rangliste", "slides": ["charts/x.png"]}, ...],
             fakta=[("Poolet TVPI", "1,40x"), ...],
             tekstslides=[{"titel": "Konklusion", "punkter": ["..."]}])
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

MASTER_NAVN = "ppim_slidemaster.pptx"

# Layout-indeks i ppim_slidemaster.pptx
L_FORSIDE = 0            # Forside_Produkt    ph 0 = titel, ph 1 = undertitel
L_AGENDA  = 2            # Agenda
L_TEKST   = 4            # Layout_basic_2     ph 0 = titel, ph 14 = indhold
L_BLANK   = 9            # Blank
L_SKILLE  = [10, 11, 12, 13]   # Skildeslide 1-4, roteres for variation
L_SLUT    = 14           # Slut-slide fondmaegler

AGENDA_PH = [34, 16, 17, 18, 19, 20, 22]   # agendaens syv raekker, oppefra og ned

PNG_RATIO = 3200 / 1800   # graferne gemmes som 3200x1800 px

# Agendaens groenne sidepanel ("Rektangel 4") er tegnet til portraetter af kontaktpersoner.
# Fladen genbruges til noegletal, saa den ikke staar tom.
PANEL       = dict(left=7.57, top=1.09, width=3.46, height=4.67)   # tommer
PANEL_TEKST = RGBColor(0xFF, 0xFF, 0xFF)
PANEL_STREG = RGBColor(0x8E, 0x94, 0x81)   # lys salvie mod den groenne flade
TITEL_FONT  = "Baskerville Old Face"       # masterens majorFont
BROED_FONT  = "Gill Sans MT"               # masterens minorFont


def _find_master(master=None):
    """Egen mappe, mappen over, ellers repo-roden (mappen over pakken)."""
    if master is not None:
        p = Path(master)
        if not p.exists():
            raise FileNotFoundError(f"Slidemasteren findes ikke: {p}")
        return p
    for p in (Path(MASTER_NAVN), Path("..") / MASTER_NAVN, Path(__file__).resolve().parents[1] / MASTER_NAVN):
        if p.exists():
            return p
    raise FileNotFoundError(f"Kunne ikke finde {MASTER_NAVN} i mappen, i '..' eller i repo-roden")


def _ryd_slides(prs):
    """Masteren indeholder seks eksempelslides - de fjernes, saa decket bygges fra bunden."""
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)


def _ny_slide(prs, layout_idx):
    return prs.slides.add_slide(prs.slide_layouts[layout_idx])


def _saet_ph(slide, idx, tekst):
    """Skriv tekst i placeholderen med det givne idx; formatering arves fra layoutet."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text_frame.text = str(tekst)
            return ph
    raise KeyError(f"Placeholder idx={idx} findes ikke paa layoutet '{slide.slide_layout.name}'")


def _fjern_tomme_ph(slide):
    """Slet klonede placeholders uden indhold - ellers staar der 'Klik for at tilfoeje ...'."""
    tree = slide.shapes._spTree
    for ph in list(slide.placeholders):
        if not (ph.has_text_frame and ph.text_frame.text.strip()):
            tree.remove(ph._element)


def _kant_til_kant(prs, slide, png):
    """Laeg billedet ud over hele sliden og skub det bagerst i tegnerækkefoelgen."""
    png = Path(png)
    if not png.exists():
        raise FileNotFoundError(f"Slide-billedet mangler - koer grafcellerne foerst: {png}")
    h = int(prs.slide_width / PNG_RATIO)
    top = int((prs.slide_height - h) / 2)      # 16:9-afrunding: under 0,01 tomme
    pic = slide.shapes.add_picture(str(png), 0, Emu(top), width=prs.slide_width, height=Emu(h))
    tree = slide.shapes._spTree
    tree.remove(pic._element)
    tree.insert(2, pic._element)               # efter nvGrpSpPr og grpSpPr
    return pic


def _fakta_panel(slide, fakta):
    """Op til tre noegletal centreret i agendaens groenne panel: label over vaerdi."""
    n = min(len(fakta), 3)
    h = PANEL["height"] / n
    for i, (label, vaerdi) in enumerate(fakta[:n]):
        top = PANEL["top"] + i * h
        if i:
            streg = slide.shapes.add_connector(
                1, Inches(PANEL["left"] + 0.45), Inches(top),
                Inches(PANEL["left"] + PANEL["width"] - 0.45), Inches(top))
            streg.line.color.rgb = PANEL_STREG
            streg.line.width = Pt(0.75)
        tf = slide.shapes.add_textbox(Inches(PANEL["left"]), Inches(top),
                                      Inches(PANEL["width"]), Inches(h)).text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        # lange vaerdier ("1.774,4 mio. USD") maa ikke sprænge panelets bredde
        vaerdi_str = 24 if len(str(vaerdi)) <= 14 else 20
        for j, (tekst, font, str_) in enumerate([(label, BROED_FONT, 10),
                                                 (vaerdi, TITEL_FONT, vaerdi_str)]):
            para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            run = para.add_run()
            run.text = str(tekst)
            run.font.name = font
            run.font.size = Pt(str_)
            run.font.color.rgb = PANEL_TEKST


def byg_deck(outfile, titel, undertitel, sektioner, fakta=None, tekstslides=None, master=None):
    """Byg decket og gem det som outfile.

    sektioner:   [{"titel": str, "slides": [sti til PNG, ...]}, ...] i praesentationsorden.
                 Hver sektion faar et skilleslide, hver PNG et fuldslide.
    fakta:       op til tre (label, vaerdi)-par til agendaens groenne panel.
    tekstslides: [{"titel": str, "punkter": [str, ...]}, ...] til sidst, foer slut-sliden.
    """
    if len(sektioner) > len(AGENDA_PH):
        raise ValueError(f"Agenda-layoutet har {len(AGENDA_PH)} raekker, men der er "
                         f"{len(sektioner)} sektioner - slaa nogle sammen.")

    prs = Presentation(str(_find_master(master)))
    _ryd_slides(prs)

    s = _ny_slide(prs, L_FORSIDE)
    _saet_ph(s, 0, titel)
    _saet_ph(s, 1, undertitel)

    s = _ny_slide(prs, L_AGENDA)
    for i, sek in enumerate(sektioner):
        _saet_ph(s, AGENDA_PH[i], f"{i + 1}.  {sek['titel']}")
    _fjern_tomme_ph(s)          # kontaktpersonernes billed- og tekstfelter til hoejre
    if fakta:
        _fakta_panel(s, fakta)

    for i, sek in enumerate(sektioner):
        s = _ny_slide(prs, L_SKILLE[i % len(L_SKILLE)])
        _saet_ph(s, 0, sek["titel"])
        for png in sek["slides"]:
            _kant_til_kant(prs, _ny_slide(prs, L_BLANK), png)

    for blok in tekstslides or []:
        s = _ny_slide(prs, L_TEKST)
        _saet_ph(s, 0, blok["titel"])
        krop = _saet_ph(s, 14, blok["punkter"][0]).text_frame
        for punkt in blok["punkter"][1:]:
            krop.add_paragraph().text = str(punkt)
        for para in krop.paragraphs:          # layoutets 12 pt er for spinkelt til en tekstslide
            para.space_before = Pt(12)
            for run in para.runs:
                run.font.size = Pt(14)

    _ny_slide(prs, L_SLUT)

    outfile = Path(outfile)
    outfile.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(outfile))
    print(f"gemt: {outfile}  ({len(prs.slides._sldIdLst)} slides)")
    return outfile
